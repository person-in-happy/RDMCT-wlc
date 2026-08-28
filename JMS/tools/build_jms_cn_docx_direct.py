"""Build the Chinese JMS review DOCX directly from the UTF-8 LaTeX master."""

from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
JMS = ROOT / "JMS"
SOURCE = JMS / "JMS_Review_CN.tex"
WORK = ROOT / ".tmp" / "jms_cn_docx_direct"
ASSETS = WORK / "assets"
OUTPUT = JMS / "output" / "docx" / "JMS_Revision_V17_CN.docx"


OUTPUT = JMS / 'output' / 'docx' / 'JMS_Revision_V18_CN.docx'


def run(*args: str, cwd: Path | None = None) -> None:
    result = subprocess.run(args, cwd=cwd or ROOT, stdout=subprocess.PIPE,
                            stderr=subprocess.STDOUT, text=True,
                            encoding="utf-8", errors="replace", check=False)
    if result.returncode:
        raise RuntimeError("Command failed: " + " ".join(args) + "\n" + result.stdout[-6000:])


def prepare() -> None:
    resolved = WORK.resolve()
    if resolved.parent != (ROOT / ".tmp").resolve():
        raise RuntimeError(f"Unsafe work path: {resolved}")
    if resolved.exists():
        shutil.rmtree(resolved)
    ASSETS.mkdir(parents=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)


def expand_inputs(text: str, base: Path) -> str:
    pattern = re.compile(r"\\input\{([^}]+)\}")
    while True:
        match = pattern.search(text)
        if not match:
            return text
        child = Path(match.group(1))
        if child.suffix != ".tex":
            child = child.with_suffix(".tex")
        if not child.is_absolute():
            child = base / child
        replacement = expand_inputs(child.read_text(encoding="utf-8"), child.parent)
        text = text[:match.start()] + replacement + text[match.end():]


def strip_comments(text: str) -> str:
    return re.sub(r"(?m)(?<!\\)%.*$", "", text)


def extract_braced(text: str, command: str) -> str:
    start = text.find(command)
    if start < 0:
        return ""
    brace = text.find("{", start + len(command))
    depth = 0
    for index in range(brace, len(text)):
        if text[index] == "{":
            depth += 1
        elif text[index] == "}":
            depth -= 1
            if depth == 0:
                return text[brace + 1:index]
    return ""


def bibliography_labels(source: str) -> dict[str, str]:
    labels = {}
    for match in re.finditer(r"\\bibitem\[(.*?)\]\{([^}]+)\}", source, re.S):
        raw = re.sub(r"\s+", " ", match.group(1).replace("~", " ")).strip()
        author_year = re.match(r"(.*?\(\d{4}[a-z]?\))", raw)
        labels[match.group(2)] = (author_year.group(1) if author_year else match.group(2)).replace(".(", ". (")
    return labels


def bibliography_labels(source: str) -> dict[str, dict[str, object]]:
    labels = {}
    pattern = re.compile(r'\\bibitem\[(.*?)\]\{([^}]+)\}', re.S)
    for number, match in enumerate(pattern.finditer(source), 1):
        raw = re.sub(
            r'\s+', ' ', match.group(1).replace('~', ' ')
        ).strip()
        author = raw.split('(', 1)[0].strip().replace(' et al.', '等')
        labels[match.group(2)] = {
            'number': number,
            'author': author,
        }
    return labels


def label_maps(source: str):
    eq, fig, tab = {}, {}, {}
    for number, block in enumerate(re.findall(r"\\begin\{equation\}.*?\\end\{equation\}", source, re.S), 1):
        match = re.search(r"\\label\{([^}]+)\}", block)
        if match:
            eq[match.group(1)] = number
    for number, block in enumerate(re.findall(r"\\begin\{figure\*?\}.*?\\end\{figure\*?\}", source, re.S), 1):
        match = re.search(r"\\label\{([^}]+)\}", block)
        if match:
            fig[match.group(1)] = number
    blocks = re.findall(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", source, re.S)
    for number, (_, block) in enumerate(sorted((source.find(block), block) for block in blocks), 1):
        match = re.search(r"\\label\{([^}]+)\}", block)
        if match:
            tab[match.group(1)] = number
    return eq, fig, tab


def unwrap_commands(text: str) -> str:
    pattern = re.compile(r"\\(?:textbf|emph|textit|texttt|mathrm|operatorname|mathcal)\{([^{}]*)\}")
    previous = None
    while text != previous:
        previous = text
        text = pattern.sub(r"\1", text)
    return text


def plain_math(text: str) -> str:
    replacements = {r"\times": "×", r"\leq": "≤", r"\le": "≤", r"\geq": "≥",
                    r"\ge": "≥", r"\in": "∈", r"\cup": "∪", r"\sum": "Σ",
                    r"\prod": "Π", r"\lambda": "λ", r"\varepsilon": "ε",
                    r"\theta": "θ", r"\alpha": "α", r"\eta": "η", r"\Delta": "Δ",
                    r"\rho": "ρ", r"\tau": "τ", r"\pi": "π", r"\nabla": "∇",
                    r"\ldots": "…", r"\cdot": "·", r"\mid": "|", r"\leftarrow": "←"}
    text = unwrap_commands(text.replace(r"\cmax", r"C_{\max}").replace(r"\pdi", "PDI"))
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\(?:bigl|bigr|Bigl|Bigr|left|right|qquad|quad)\b", " ", text)
    text = re.sub(r"\\([A-Za-z]+)", r"\1", text.replace("{", "").replace("}", ""))
    return re.sub(r"\s+", " ", text).strip()


def make_plain_converter(citations, eq, fig, tab):
    def citep(match):
        labels = [citations.get(key.strip(), key.strip()) for key in match.group(1).split(",")]
        return "（" + "；".join(labels) + "）"

    def citet(match):
        return "、".join(citations.get(key.strip(), key.strip()) for key in match.group(1).split(","))

    def numeric_group(keys):
        numbers = sorted({
            int(citations[key]['number'])
            for key in keys if key in citations
        })
        groups = []
        start = end = None
        for number in numbers:
            if start is None:
                start = end = number
            elif number == end + 1:
                end = number
            else:
                groups.append((start, end))
                start = end = number
        if start is not None:
            groups.append((start, end))
        parts = []
        for start, end in groups:
            if end - start >= 2:
                parts.append(f'{start}–{end}')
            elif end == start + 1:
                parts.extend((str(start), str(end)))
            else:
                parts.append(str(start))
        return '[' + ','.join(parts) + ']'

    def citep(match):
        keys = [key.strip() for key in match.group(1).split(',')]
        return numeric_group(keys)

    def citet(match):
        values = []
        for key in (item.strip() for item in match.group(1).split(',')):
            record = citations.get(key)
            values.append(
                f'{record["author"]}[{record["number"]}]'
                if record else key
            )
        return '、'.join(values)

    def convert(text: str) -> str:
        text = text.replace(r'\DOIprefix', ' ').replace(r'\URLprefix', ' ')
        text = re.sub(r"\\citep\{([^}]+)\}", citep, text)
        text = re.sub(r"\\citet\{([^}]+)\}", citet, text)
        text = re.sub(r"\\eqref\{([^}]+)\}", lambda m: f"（{eq.get(m.group(1), '?')}）", text)
        refs = {**fig, **tab, **eq}
        text = re.sub(r"\\ref\{([^}]+)\}", lambda m: str(refs.get(m.group(1), "?")), text)
        text = re.sub(r"\\label\{[^}]+\}", "", text)
        text = text.replace(r"\method{}", "RDMCT-HBS").replace(r"\method", "RDMCT-HBS")
        text = text.replace(r"\cmax", r"$C_{\max}$").replace(r"\pdi", "PDI")
        text = re.sub(r"\$([^$]+)\$", lambda m: plain_math(m.group(1)), text)
        text = re.sub(r"\\url\{([^}]+)\}", r"\1", text)
        text = re.sub(r"\\doi\{([^}]+)\}", r"doi: \1", text)
        text = unwrap_commands(text)
        text = text.replace(r"\%", "%").replace(r"\&", "&").replace(r"\_", "_")
        text = text.replace("~", " ").replace("--", "–")
        text = re.sub(r"\\(?:noindent|small|normalsize|centering|newpage)\b", "", text)
        text = re.sub(r"\\[A-Za-z@]+(?:\[[^]]*\])?", "", text)
        text = text.replace("{", "").replace("}", "").replace("\\", " ")
        return re.sub(r"\s+", " ", text).strip()

    return convert


def render_equation(block: str, number: int) -> Path:
    clean = re.sub(r"\\label\{[^}]+\}", "", block)
    clean = re.sub(r"\n\s*\n", "\n", clean)
    wrapper = "\n".join([
        r"\documentclass[varwidth,border=5pt]{standalone}",
        r"\usepackage[UTF8,scheme=plain,fontset=windows]{ctex}",
        r"\usepackage{amsmath,amssymb,mathtools}",
        r"\begin{document}", rf"\setcounter{{equation}}{{{number - 1}}}",
        clean, r"\end{document}"])
    tex = WORK / f"equation_{number:02d}.tex"
    tex.write_text(wrapper, encoding="utf-8")
    run("xelatex", "-interaction=nonstopmode", "-halt-on-error", f"-output-directory={WORK}", str(tex), cwd=WORK)
    output = ASSETS / f"equation_{number:02d}"
    run("pdftoppm", "-png", "-r", "300", "-singlefile", str(WORK / f"equation_{number:02d}.pdf"), str(output))
    return output.with_suffix(".png")


def render_figure(name: str, number: int) -> Path:
    source = JMS / "figures" / f"{name}.pdf"
    if not source.exists():
        raise FileNotFoundError(source)
    output = ASSETS / f"figure_{number:02d}"
    run("pdftoppm", "-png", "-r", "240", "-singlefile", str(source), str(output))
    return output.with_suffix(".png")


def add_picture_paragraph(
    doc: Document,
    path: Path,
    max_width_cm: float,
    keep_next: bool = False,
    alt_text: str = "",
) -> None:
    with Image.open(path) as image:
        natural_cm = image.width / 300 * 2.54
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = keep_next
    shape = paragraph.add_run().add_picture(
        str(path), width=Cm(min(max_width_cm, max(5.0, natural_cm)))
    )
    if alt_text:
        shape._inline.docPr.set("descr", alt_text)
        shape._inline.docPr.set("title", alt_text)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_obj = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run_obj._r.extend([begin, instr, end])


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run_obj = cell.paragraphs[0].add_run(text)
    run_obj.bold = bold
    run_obj.font.name = "Times New Roman"
    run_obj._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run_obj.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def parse_table_rows(block: str, convert) -> list[list[str]]:
    rows = []
    for raw in block.splitlines():
        line = raw.strip()
        if not line.endswith(r"\\") or r"\multicolumn" in line:
            continue
        content = line[:-2]
        cells = [convert(part) for part in (content.split("&") if "&" in content else [content])]
        if not cells or any(not cell for cell in cells) or (rows and cells == rows[0]):
            continue
        rows.append(cells)
    return rows


def add_body_text(doc: Document, text: str, convert) -> None:
    text = re.sub(r"\\begin\{itemize\}.*?\\end\{itemize\}", "", text, flags=re.S)
    for chunk in re.split(r"\n\s*\n", text):
        value = convert(chunk.replace("\n", " "))
        if not value:
            continue
        paragraph = doc.add_paragraph(value, style="Body Text")
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.2


def parse_bibliography(source: str, convert) -> list[str]:
    env = re.search(r"\\begin\{thebibliography\}.*?\n(.*?)\\end\{thebibliography\}", source, re.S)
    if not env:
        return []
    body = env.group(1)
    markers = list(re.finditer(r"\\bibitem(?:\[.*?\])?\{[^}]+\}", body, re.S))
    entries = []
    for index, marker in enumerate(markers):
        end = markers[index + 1].start() if index + 1 < len(markers) else len(body)
        value = convert(body[marker.end():end].replace("\n", " "))
        raw_entry = body[marker.end():end].replace('\n', ' ')
        raw_entry = re.sub(r'\\bibinfo\{[^}]+\}', '', raw_entry)
        value = convert(raw_entry)
        if value:
            entries.append(value)
    return [f'[{index}] {entry}' for index, entry in enumerate(entries, 1)]


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin = section.right_margin = Cm(2.2)
    add_page_number(section)
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(10.5)
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size, style.font.bold = Pt(size), True


def build() -> None:
    source = strip_comments(expand_inputs(SOURCE.read_text(encoding="utf-8"), JMS))
    citations = bibliography_labels(source)
    eq_labels, fig_labels, tab_labels = label_maps(source)
    convert = make_plain_converter(citations, eq_labels, fig_labels, tab_labels)
    title = convert(extract_braced(source, r"\title"))
    abstract_match = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", source, re.S)
    keywords_match = re.search(r"\\begin\{keyword\}(.*?)\\end\{keyword\}", source, re.S)
    if not abstract_match or not keywords_match:
        raise RuntimeError("Front matter not found")

    document = Document()
    style_document(document)
    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run(title)
    for value in (
        "邵立珍（Lizhen Shao）*；魏林聪（Lincong Wei）；王蔚（Wei Wang）",
        "北京科技大学控制科学与工程系，北京市海淀区学院路30号，北京100083，中国",
        "*通讯作者：lshao@ustb.edu.cn",
    ):
        paragraph = document.add_paragraph(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("摘要", level=1)
    add_body_text(document, abstract_match.group(1), convert)
    keyword_p = document.add_paragraph()
    keyword_p.add_run("关键词：").bold = True
    keyword_p.add_run(convert(keywords_match.group(1).replace(r"\sep", "；")))

    body_match = re.search(r"(\\section\{引言\}.*?)(?=\\begin\{thebibliography\})", source, re.S)
    if not body_match:
        raise RuntimeError("Manuscript body not found")
    body = body_match.group(1)
    block_pattern = re.compile(
        r"(?P<section>\\section\*?\{[^}]+\})|(?P<subsection>\\subsection\{[^}]+\})|"
        r"(?P<figure>\\begin\{figure\*?\}.*?\\end\{figure\*?\})|"
        r"(?P<table>\\begin\{table\*?\}.*?\\end\{table\*?\})|"
        r"(?P<equation>\\begin\{equation\}.*?\\end\{equation\})", re.S)
    cursor = equation_number = figure_number = table_number = algorithm_number = 0
    for match in block_pattern.finditer(body):
        add_body_text(document, body[cursor:match.start()], convert)
        block = match.group(0)
        if match.lastgroup == "section":
            document.add_heading(convert(extract_braced(block, r"\section")), level=1)
        elif match.lastgroup == "subsection":
            document.add_heading(convert(extract_braced(block, r"\subsection")), level=2)
        elif match.lastgroup == "equation":
            equation_number += 1
            add_picture_paragraph(
                document,
                render_equation(block, equation_number),
                15.4,
                alt_text=f"公式 {equation_number}",
            )
        elif match.lastgroup == "figure":
            figure_number += 1
            graphic = re.search(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", block)
            if not graphic:
                raise RuntimeError(f"Figure {figure_number} has no external graphic")
            add_picture_paragraph(
                document,
                render_figure(graphic.group(1), figure_number),
                15.6,
                True,
                alt_text=f"图 {figure_number}",
            )
            caption = document.add_paragraph(
                f"图 {figure_number}. {convert(extract_braced(block, r'\caption'))}", style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_previous = True
        else:
            table_number += 1
            algorithm_title = extract_braced(block, r"\algorithmtitle")
            if algorithm_title:
                algorithm_number += 1
                caption_text = f"算法 {algorithm_number}：{convert(algorithm_title)}"
            else:
                caption_text = f"表 {table_number}. {convert(extract_braced(block, r'\caption'))}"
            caption = document.add_paragraph(caption_text, style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_next = True
            rows = parse_table_rows(block, convert)
            if not rows:
                raise RuntimeError(f"Table {table_number} has no parsed rows")
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            header_properties = table.rows[0]._tr.get_or_add_trPr()
            header_marker = OxmlElement("w:tblHeader")
            header_marker.set(qn("w:val"), "true")
            header_properties.append(header_marker)
            for row_index, row in enumerate(rows):
                for col_index in range(width):
                    value = row[col_index] if col_index < len(row) else ""
                    set_cell_text(table.cell(row_index, col_index), value, row_index == 0)
        cursor = match.end()
    add_body_text(document, body[cursor:], convert)

    document.add_heading("参考文献", level=1)
    references = parse_bibliography(source, convert)
    for reference in references:
        paragraph = document.add_paragraph(reference, style="Body Text")
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)
        paragraph.paragraph_format.space_after = Pt(3)
    if (equation_number, figure_number, table_number, len(references)) != (34, 11, 5, 30):
        raise RuntimeError(
            f"Unexpected counts: equations={equation_number}, figures={figure_number}, "
            f"tables={table_number}, references={len(references)}")

    document.core_properties.title = title
    document.core_properties.subject = "Journal of Manufacturing Systems 中文审阅稿"
    document.core_properties.author = "邵立珍；魏林聪；王蔚"
    temp = WORK / "JMS_Revision_V17_CN.docx"
    temp = WORK / 'JMS_Revision_V18_CN.docx'
    document.save(temp)
    check = Document(temp)
    all_text = "\n".join(paragraph.text for paragraph in check.paragraphs)
    required = ("工业问题与两阶段MILP", "符号与紧凑模型", "规模扩展与工业边界", "公式 （13）表示")
    required = (
        '工业问题与两阶段MILP',
        '建模边界、物料路线与符号',
        '核心公式的覆盖范围与模型边界',
        '公式 （25）表示',
        '规模、混流构成与求解难度',
        '分布偏移下的可行排程交付',
        '工业解释与综合结论',
        'CT调度研究首先围绕吞吐分析和驻留时间可行性展开',
        '两阶段工作流使路线总等待、最大等待、节拍变异系数和节拍偏差分别降低',
        '设备可配置多个旋转腔室',
        'ATR先把晶圆放至AL完成对准',
        '已有学习型割选择分别以强化学习',
        'SPBS按实现中的五级配置逐步扩大搜索空间',
        '从候选池到有序割集的逐步选择',
        '步骤5：束搜索解码',
        '角色次模补全',
        '基于并行层次REINFORCE与束搜索的割平面选择算法',
        '每个含产品PW对必须连续出现在相邻两个加工周期中',
        'PW对（wafer pair）是由两个晶圆组成的物理搬运与腔室占位单元',
        '强化学习训练与精确求解接口',
        '把HEM的13维实现作为直接基线',
        '魏林聪负责概念设计、方法设计、软件实现',
        '作者使用OpenAI ChatGPT和Codex改进语言和可读性',
    )
    if any(item not in all_text for item in required):
        raise RuntimeError("Rebuilt DOCX is missing revised model content")
    if "�" in all_text:
        raise RuntimeError("Unicode replacement characters found in rebuilt DOCX")
    shutil.copy2(temp, OUTPUT)
    print(OUTPUT)
    print(f"paragraphs={len(check.paragraphs)} tables={len(check.tables)} "
          f"shapes={len(check.inline_shapes)} refs={len(references)}")


if __name__ == "__main__":
    prepare()
    build()
